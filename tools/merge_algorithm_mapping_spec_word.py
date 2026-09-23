from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "software_stack_spec.zh-CN.md"
OUTPUT = ROOT / "docs" / "software_stack_spec.zh-CN.docx"
STAGED_OUTPUT = ROOT / "docs" / "software_stack_spec.zh-CN.merged.docx"
START = "### 10.1 当前算子支持矩阵"
END = "## 11 `.ftlpu` 可执行程序格式与 loader"


def element_text(element) -> str:
    return "".join(
        child.text or ""
        for child in element.iter(qn("w:t"))
    )


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, width in (("top", 65), ("bottom", 65), ("left", 80), ("right", 80)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")


def add_table_before(doc: Document, anchor, rows: list[list[str]]) -> None:
    width_sets = {
        3: (1.75, 2.55, 2.70),
    }
    widths = width_sets.get(len(rows[0]), tuple(7.0 / len(rows[0]) for _ in rows[0]))
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "1D344A")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F7FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(7.5 if row_index else 7.7)
            run.bold = row_index == 0
            if row_index == 0:
                run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            header_flag = OxmlElement("w:tblHeader")
            header_flag.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(header_flag)
    anchor.addprevious(table._tbl)


def add_paragraph_before(doc: Document, anchor, value: str, heading: bool) -> None:
    paragraph = doc.add_paragraph(style="Heading 3" if heading else "Normal")
    paragraph.paragraph_format.keep_with_next = heading
    run = paragraph.add_run(value)
    run.font.name = "Microsoft YaHei"
    if heading:
        run.font.size = Pt(10.5)
        run.bold = True
        run.font.color.rgb = RGBColor(29, 52, 74)
    else:
        run.font.size = Pt(9)
    anchor.addprevious(paragraph._p)


def parse_rows(lines: list[str]) -> list[list[str]]:
    rows = [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in lines
    ]
    return [
        row
        for row in rows
        if not all(set(cell) <= {"-", ":", " "} for cell in row)
    ]


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    start = source.index(START)
    end = source.index(END, start)
    lines = source[start:end].splitlines()
    doc = Document(OUTPUT)
    body = doc.element.body
    anchor = next(
        element
        for element in body
        if element.tag == qn("w:p")
        and element_text(element).startswith("11 `.ftlpu`")
    )
    previous_start = next(
        (
            element
            for element in body
            if element.tag == qn("w:p")
            and element_text(element).startswith("10.1 当前算子支持矩阵")
        ),
        None,
    )
    if previous_start is not None:
        current = previous_start
        while current is not anchor:
            successor = current.getnext()
            body.remove(current)
            current = successor

    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("### "):
            add_paragraph_before(doc, anchor, line[4:], heading=True)
            index += 1
            continue
        if line.startswith("|"):
            group = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                group.append(lines[index])
                index += 1
            add_table_before(doc, anchor, parse_rows(group))
            continue
        add_paragraph_before(doc, anchor, line, heading=False)
        index += 1

    version_line = source.splitlines()[2]
    version_paragraph = next(
        paragraph for paragraph in doc.paragraphs
        if paragraph.text.startswith("版本：")
    )
    version_paragraph.runs[0].text = version_line
    regression = next(
        table for table in doc.tables
        if table.rows[0].cells[0].text == "改动类型"
    )
    row = next(
        row for row in regression.rows
        if row.cells[0].text in {"Qwen decoder 改动", "Decoder 算子映射改动"}
    )
    replacements = (
        "Decoder 算子映射改动",
        "Qwen2.5/Qwen3 或 SmolLM2 对应单层 pipeline、数值与 KV state 测试",
        "Attention/FFN 分阶段误差、shape/type 拒绝、实际 C2C 重叠",
    )
    for cell, value in zip(row.cells, replacements):
        cell.text = value
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(7)

    doc.core_properties.subject = (
        "FTLPU software stack, operator support, algorithm mapping, "
        "runtime, C2C, and ICU instructions"
    )
    doc.save(STAGED_OUTPUT)
    STAGED_OUTPUT.replace(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
